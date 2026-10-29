#!/usr/bin/env python3
"""Auditoria numérica do DOCX contra as tabelas de origem."""
from pathlib import Path
import re
import numpy as np
import pandas as pd
import docx

ROOT = Path(__file__).resolve().parents[1]
DOCX = ROOT / "artigo_final" / "artigo_expressao_diferencial_covariaveis_deconvolucao.docx"
TAB = ROOT / "resultados" / "analises_avancadas" / "tabelas"
DE = TAB / "expressao_diferencial_covariaveis"
OUT = ROOT / "auditorias"
OUT.mkdir(exist_ok=True)

d = docx.Document(DOCX)
paras = [p.text for p in d.paragraphs]
tables = ["\n".join(c.text for r in t.rows for c in r.cells) for t in d.tables]
body = "\n".join(paras)
alltext = body + "\n" + "\n".join(tables)


def br(x, dec=0):
    s = f"{x:,.{dec}f}"
    return s.replace(",", "X").replace(".", ",").replace("X", ".")


def has(s):
    return s in alltext


def has_par(s):
    return s in body


checks = []


def ck(label, ok, expected, source):
    checks.append(dict(verificacao=label, resultado="PASS" if ok else "FAIL",
                       esperado=str(expected), fonte=source))


# ---------------------------------------------------------------------------
# Coorte e dados
# ---------------------------------------------------------------------------
meta = pd.read_csv(ROOT / "dados" / "processed" / "metadata_analise.tsv", sep="\t", low_memory=False)
meta = meta[~((meta["group"] == "GTEx_Normal") & meta["subject"].duplicated())]
n = meta["group"].value_counts()
ck("Total de amostras", has("528") and len(meta) == 528, 528, "metadata_analise.tsv")
ck("Tumores", has("369") and n.get("TCGA_Tumor") == 369, 369, "metadata_analise.tsv")
ck("Adjacentes", has("50") and n.get("TCGA_Adjacent") == 50, 50, "metadata_analise.tsv")
ck("GTEx", has("109") and n.get("GTEx_Normal") == 109, 109, "metadata_analise.tsv")
ck("Genes", has("12.401"), 12401, "matriz de expressão")
ck("Amostras antes da deduplicação", has("529"), 529, "metadata bruto")

tum = meta[meta["group"] == "TCGA_Tumor"]
for var, col, val in [("invasão vascular", "vascular_invasion", 70.7), ("sorologia", "viral_status", 55.8),
                      ("fibrose", "fibrosis_ishak", 43.1), ("Child-Pugh", "child_pugh", 35.8), ("AFP", "afp", 25.2)]:
    if col == "viral_status":
        pct = 100 * (tum[col].isna() | (tum[col].astype(str).str.strip().isin(["", "Unknown"]))).mean()
    else:
        pct = 100 * (tum[col].isna() | (tum[col].astype(str).str.strip() == "")).mean()
    ck(f"Missingness {var}", abs(pct - val) < 0.1 and has(br(val, 1).replace(".0", "")), br(val, 1), col)

ck("Child-Pugh ausentes", has("132") and (tum["child_pugh"].isna()).sum() == 132, 132, "metadata")
ck("Child-Pugh B-C corrigido", has("22") and ((tum["child_pugh"].isin(["B", "C"])).sum()) == 22, 22, "metadata")
ck("Child-Pugh código antigo", has("154"), 154, "132 ausentes + 22 B-C")
ck("Child-Pugh GDC Unknown", has("84"), 84, "auditoria_child_pugh_gdc.md")
ck("Child-Pugh GDC vazios", has("48"), 48, "auditoria_child_pugh_gdc.md")
ck("Concordância GDC", has("100%") and has("237"), 237, "auditoria_child_pugh_gdc.md")
ck("GDC estágio/grau/fibrose", has("345") and has("364") and has("210"), "345/364/210", "auditoria_gdc_outros.csv")

# ---------------------------------------------------------------------------
# Expressão diferencial
# ---------------------------------------------------------------------------
res = pd.read_csv(DE / "resumo_modelos_expressao_diferencial_covariaveis.csv").set_index("modelo")
for mod, val in [("tumor_vs_adjacente_base", 4130), ("tumor_vs_adjacente_pareado", 3710),
                 ("tumor_vs_adjacente_clinico_completo", 3740), ("tumor_vs_adjacente_clinico_tss", 3736),
                 ("estagio_base", 46), ("estagio_base_mesmos_casos", 0),
                 ("grau_base", 172), ("grau_clinico", 199), ("grau_clinico_tss", 247),
                 ("grau_so_tss_full", 337), ("grau_base_sem_estagio", 153), ("grau_sem_DD", 276),
                 ("sexo_gtex", 13), ("sexo_tumor", 175), ("sexo_adjacente", 103), ("afp_400", 455),
                 ("tumor_vs_gtex_base", 3391)]:
    ck(f"DEGs {mod}", int(res.loc[mod, "total_degs"]) == val and has(br(val)), val, "resumo_modelos...csv")
ck("Número de modelos", len(res) == 48 and has("48"), 48, "resumo_modelos...csv")

comp = pd.read_csv(TAB / "integracao_composicao" / "comparacao_base_vs_composicao.csv").set_index("comparacao")
for mod, val in [("tumor_vs_adjacente_composicao", 1897), ("tumor_vs_gtex_composicao", 2732),
                 ("estagio_composicao", 0), ("grau_composicao", 0),
                 ("sexo_gtex_composicao", 15), ("sexo_tumor_clinico_composicao", 90),
                 ("sexo_adjacente_composicao", 22)]:
    if mod in comp.index:
        ck(f"DEGs composição {mod}", int(comp.loc[mod, "degs_ajustado"]) == val and has(br(val)), val,
           "comparacao_base_vs_composicao.csv")

# ---------------------------------------------------------------------------
# ComBat
# ---------------------------------------------------------------------------
cb = pd.read_csv(TAB / "resumo_testes_ComBat.csv").set_index("cenario")
ck("Cenários ComBat", len(cb) == 17 and has("17"), 17, "resumo_testes_ComBat.csv")
for c, antes, depois in [("tcga_gtex_sem_mod", 3391, 1), ("tcga_gtex_sem_mod_meanonly", 3391, 1947),
                         ("tcga_tss_grau", 172, 290), ("tcga_tss_estagio", 46, 0),
                         ("tcga_ta_tss_mod_sexo_idade", 4130, 3048),
                         ("tcga_ta_tss_sem_mod", 4130, 3038),
                         ("tcga_ta_tss_sem_mod_meanonly", 4130, 3605)]:
    ok = int(cb.loc[c, "degs_antes"]) == antes and int(cb.loc[c, "degs_depois"]) == depois
    ck(f"ComBat {c}", ok and has(br(antes)) and has(br(depois)), f"{antes}→{depois}", "resumo_testes_ComBat.csv")

# ---------------------------------------------------------------------------
# Deconvolução
# ---------------------------------------------------------------------------
cg = pd.read_csv(TAB / "deconvolucao" / "comparacao_celulas_por_grupo.csv").set_index("celula")
for cel, vals in [("xCell_Hepatocytes", [0.151, 0.225, 0.206]),
                  ("xCell_Endothelial cells", [0.0406, 0.0222, 1.99e-18]),
                  ("xCell_Macrophages", [0.0161, 0.049, 0.0181])]:
    row = cg.loc[cel]
    got = [row["mediana_tumor"], row["mediana_adjacente"], row["mediana_gtex"]]
    ck(f"Medianas {cel}", all(abs(a - b) < max(0.002, abs(b) * 0.05) for a, b in zip(vals, got)),
       vals, "comparacao_celulas_por_grupo.csv")
ck("Concordância mediana 0,01", has("0,01"), 0.01, "concordancia_metodos_deconvolucao.csv")
for s in ["−0,33", "−0,38", "−0,51", "−0,31", "0,47", "0,62", "0,64"]:
    ck(f"Correlação deconvolução {s}", has(s), s, "QC_correlacao_celulas_marcadores.csv")

# ---------------------------------------------------------------------------
# Cox corrigido
# ---------------------------------------------------------------------------
cox = pd.read_csv(TAB / "cox_clinico_corrigido.csv").set_index("term")
ck("Cox n", int(cox.loc["stage_groupIII-IV", "n"]) == 337 and has("337"), 337, "cox_clinico_corrigido.csv")
ck("Cox eventos", int(cox.loc["stage_groupIII-IV", "eventos"]) == 112 and has("112"), 112, "cox_clinico_corrigido.csv")
ck("Cox estágio HR", abs(cox.loc["stage_groupIII-IV", "estimate"] - 2.61418) < 0.001 and has("2,61"), 2.61,
   "cox_clinico_corrigido.csv")
ck("Cox grau HR", abs(cox.loc["grade_groupG3-G4", "estimate"] - 1.159997) < 0.001 and has("1,16"), 1.16,
   "cox_clinico_corrigido.csv")
ck("Log-rank", has("3,2 × 10⁻⁷"), "3,2e-7", "logrank_estagio.csv")
ck("C-index", has("0,637") and has("0,027"), "0,637", "concordance(fitA)")
zph = pd.read_csv(TAB / "cox_zph.csv")
ck("cox.zph global clínico", has("0,386"), 0.386, "cox_zph.csv")
ck("cox.zph global AFP/fibrose", has("0,153"), 0.153, "cox_zph.csv")

# ---------------------------------------------------------------------------
# Sensibilidades, GDC, validação externa
# ---------------------------------------------------------------------------
sc = pd.read_csv(TAB / "sensibilidades" / "sensibilidade_composicao_resumo.csv")
for cen in ["xCell_sem_marcadores", "xCell_sem_correlatos", "xCell_universo_restrito"]:
    ck(f"Composição {cen} estágio 0", int(sc[(sc.contraste == "estagio") & (sc.cenario == cen)].degs.iloc[0]) == 0, 0,
       "sensibilidade_composicao_resumo.csv")
ck("Composição escore curado grau 27", int(sc[(sc.contraste == "grau") & (sc.cenario == "escore_curado")].degs.iloc[0]) == 27 and has("27"),
   27, "sensibilidade_composicao_resumo.csv")
ck("Genes removidos por rho", (sc.cenario == "xCell_sem_correlatos").any() and has("279"), 279,
   "sensibilidade_composicao_resumo.csv")

ipw = pd.read_csv(TAB / "sensibilidades" / "sensibilidade_missingness_ipw.csv").set_index("contraste")
for k, auc in [("sorologia_hbv_vs_hcv", 0.808), ("child_pugh", 0.927), ("fibrose", 0.854), ("invasao_vascular", 0.729)]:
    ck(f"IPW AUC {k}", abs(ipw.loc[k, "auc_modelo_observacao"] - auc) < 0.001 and has(br(auc, 3)), br(auc, 3),
       "sensibilidade_missingness_ipw.csv")
ck("IPW Child-Pugh 1 DEG", int(ipw.loc["child_pugh", "degs_com_peso"]) == 1 and has("H2BC4"), 1,
   "sensibilidade_missingness_ipw.csv")

vm = pd.read_csv(TAB / "validacao_externa" / "validacao_multicoorte.csv").set_index(["contraste", "coorte"])
for coh, degs, com, jac, auc in [("GSE14520", 1061, 743, 0.167, 0.989), ("GSE36376", 667, 564, 0.133, 0.990),
                                 ("GSE76427", 509, 358, 0.084, 0.957)]:
    ck(f"Tumor/adj {coh}", int(vm.loc[("Tumor vs adjacente", coh), "degs_coorte"]) == degs and has(br(degs)),
       degs, "validacao_multicoorte.csv")
    ck(f"Tumor/adj {coh} sobreposição", int(vm.loc[("Tumor vs adjacente", coh), "degs_comuns"]) == com and has(br(com)),
       com, "validacao_multicoorte.csv")
    ck(f"Tumor/adj {coh} AUC", abs(vm.loc[("Tumor vs adjacente", coh), "auc_tcga"] - auc) < 0.002 and has(br(auc, 3)),
       br(auc, 3), "validacao_multicoorte.csv")
for coh, degs, auc in [("GSE14520", 76, 0.757), ("GSE36376", 0, 0.467), ("GSE76427", 0, 0.626)]:
    ck(f"Estágio {coh}", int(vm.loc[("Estágio III-IV vs I-II", coh), "degs_coorte"]) == degs,
       degs, "validacao_multicoorte.csv")
    ck(f"Estágio {coh} AUC", abs(vm.loc[("Estágio III-IV vs I-II", coh), "auc_tcga"] - auc) < 0.002 and has(br(auc, 3)),
       br(auc, 3), "validacao_multicoorte.csv")
ck("Grau GSE62232 0 DEGs", int(vm.loc[("Grau G3-G4 vs G1-G2", "GSE62232"), "degs_coorte"]) == 0, 0,
   "validacao_multicoorte.csv")
ck("Grau GSE62232 AUC", abs(vm.loc[("Grau G3-G4 vs G1-G2", "GSE62232"), "auc_tcga"] - 0.655) < 0.002 and has("0,655"),
   0.655, "validacao_multicoorte.csv")
for coh, degs, auc in [("GSE14520", 5, 0.633), ("GSE76427", 5, 0.665), ("GSE62232", 9, 0.748)]:
    ck(f"Sexo {coh}", int(vm.loc[("Masculino vs feminino", coh), "degs_coorte"]) == degs,
       degs, "validacao_multicoorte.csv")
    ck(f"Sexo {coh} AUC", abs(vm.loc[("Masculino vs feminino", coh), "auc_tcga"] - auc) < 0.002 and has(br(auc, 3)),
       br(auc, 3), "validacao_multicoorte.csv")
ck("Estágio 17/29", has("17 dos 29"), "17/29", "validacao_multicoorte.csv")
cvs = pd.read_csv(TAB / "validacao_externa" / "validacao_cruzada_grau_resumo.csv").iloc[0]
ck("CV grau AUC", abs(cvs.auc_media - 0.7047) < 0.002 and has("0,705"), 0.705, "validacao_cruzada_grau_resumo.csv")
ck("LOCO mínimo/máximo", has("0,720") and has("0,843"), "0,720–0,843", "validacao_leave_one_center_out_grau.csv")

ms = pd.read_csv(TAB / "sensibilidades" / "modelos_mistos_resumo.csv").set_index("contraste")
ck("Misto estágio 0", int(ms.loc["estagio", "degs_duplicateCorrelation"]) == 0, 0, "modelos_mistos_resumo.csv")
ck("Misto grau 283", int(ms.loc["grau", "degs_duplicateCorrelation"]) == 283 and has("283"), 283,
   "modelos_mistos_resumo.csv")

# ---------------------------------------------------------------------------
# Contagens (se disponíveis)
# ---------------------------------------------------------------------------
cont_path = TAB / "sensibilidades" / "contagens_resumo_consolidado.csv"
if cont_path.exists():
    cr = pd.read_csv(cont_path)
    for _, r in cr.iterrows():
        val = int(r["degs"])
        ck(f"Contagens {r['universo']} {r['metodo']} {r['contraste']}", has(br(val)), val,
           "contagens_resumo_consolidado.csv")
    for v in [0.325, 0.051, 0.152, 0.913]:
        ck(f"Contagens métrica {v}", has(br(v, 3)), br(v, 3), "contagens_resumo_consolidado.csv")
    ck("Contagens comum estágio 4 genes", has("4 genes em comum"), 4, "contagens_resumo_consolidado.csv")
    ck("Contagens comum grau 32 genes", has("32 genes em comum"), 32, "contagens_resumo_consolidado.csv")
else:
    ck("Contagens DESeq2/edgeR", False, "arquivo ausente", "contagens_resumo_consolidado.csv")

# ---------------------------------------------------------------------------
# Estrutura, estilo e referências
# ---------------------------------------------------------------------------
n_fig = len(d.inline_shapes)
n_tab = len(d.tables)
caps = [p for p in paras if re.match(r"^(Figura|Tabela|Quadro) (S?\d+)\.", p)]
ck("Figuras", n_fig >= 10, "10+ (5 ampliadas)", "DOCX")
ck("Tabelas", n_tab >= 19, "19+", "DOCX")
ck("Tabela S8 27 genes", has("Tabela S8.") and has("TROAP") and has("MYBL2"), "27 genes", "degs_grau__escore_curado.csv")
ck("Legendas com fonte", all("Fonte." in p for p in caps), len(caps), "DOCX")
ck("Sem marcadores de lista", sum(1 for p in d.paragraphs if p.style.name.startswith("List")) == 0, 0, "DOCX")
body_main = body.split("REFERÊNCIAS")[0]
ck("Sem dois pontos no corpo", ":" not in body_main, 0, "DOCX")
ck("Referências", len([p for p in paras if re.match(r"^\[\d+\]", p)]) == 35, 35, "DOCX")
ck("Datas de acesso distintas", len(set(re.findall(r"Acesso em (\d{1,2}) out", alltext))) == 14, 14, "DOCX")
ck("Figura 5 FDR 0,64", has("0,64"), 0.64, "modelo_estagio_base_mesmos_casos_completo.csv")
for old in ["2,56", "4,5 × 10⁻⁸", "186; 53", "128 eventos", "1,23 (0,86"]:
    ck(f"Número antigo ausente {old}", not has(old), "ausente", "DOCX")

# ---------------------------------------------------------------------------
df = pd.DataFrame(checks)
df.to_csv(OUT / "auditoria_numerica.csv", index=False)
n_pass = (df.resultado == "PASS").sum()
lines = ["# Auditoria numérica\n",
         f"Verificações: **{len(df)}**. PASS: **{n_pass}**. FAIL: **{len(df)-n_pass}**.\n",
         "| Verificação | Resultado | Esperado | Fonte |", "|---|---|---|---|"]
for _, r in df.iterrows():
    lines.append(f"| {r['verificacao']} | {r['resultado']} | {r['esperado']} | {r['fonte']} |")
if len(df) - n_pass:
    lines.append("\n## Falhas\n")
    for _, r in df[df.resultado == "FAIL"].iterrows():
        lines.append(f"- {r['verificacao']} (esperado {r['esperado']}, fonte {r['fonte']}).")
(OUT / "auditoria_numerica.md").write_text("\n".join(lines))
print(df["resultado"].value_counts().to_string())
print("Falhas:")
print(df[df.resultado == "FAIL"][["verificacao", "esperado"]].to_string(index=False))
print("Arquivo:", OUT / "auditoria_numerica.md")
